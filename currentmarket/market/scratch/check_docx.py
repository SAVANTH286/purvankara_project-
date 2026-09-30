import docx

doc = docx.Document('Puravankara_AI_DSS_Comprehensive_Model_Guide.docx')
print("Total paragraphs:", len(doc.paragraphs))
for i, p in enumerate(doc.paragraphs):
    runs_with_drawings = [r for r in p.runs if len(r._element.xpath('.//w:drawing')) > 0]
    if runs_with_drawings:
        print(f"Paragraph {i} has {len(runs_with_drawings)} run(s) with drawings.")
    if "Figure" in p.text:
        print(f"Paragraph {i} text: {p.text}")
